# -*- coding: utf-8 -*-
"""Независимая сверка aquaculture_report.xlsx (все 9 листов) с pandas-пересчётом CSV.
Ключевые моменты: все CSV читаются как строки (dtype=str), ID/коды сравниваются
как в БД; ISSCAAP CODE извлекается так же, как в SQL ((CODE mod 10) + 10*((CODE div 100))).
"""
import pandas as pd
from openpyxl import load_workbook

YEARS = [2020, 2021, 2022, 2023, 2024]
YC = ["VALUE_Y%d" % y for y in YEARS]

R = lambda f, **kw: pd.read_csv("CSV/%s.csv" % f, dtype=str, keep_default_na=False, **kw)

q = R("TSD_AQUACULTURE_QUANTITY")
v = R("TSD_AQUACULTURE_VALUE")
sp = R("REF_SPECIES")
ctry = R("REF_COUNTRY")
area = R("REF_AREA")
rel_ima = R("REL_INLAND_MARINE_AREA")
ima = R("REF_INLAND_MARINE")
yg = R("REL_YEARBOOK_GROUP_SPECIES")
iss_rel = R("REL_ISSCAAP_GROUP_SPECIES")
iss = R("REF_ISSCAAP_GROUP")

for df in (q, v):
    for c in YC:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)

excl = set(yg[yg.PARENT_ID == "00000004"].CHILD_ID)
g2 = set(yg[yg.PARENT_ID == "00000002"].CHILD_ID)
g3 = set(yg[yg.PARENT_ID == "00000003"].CHILD_ID)
sid = sp.drop_duplicates("ALPHA_3_CODE").set_index("ALPHA_3_CODE").ID
_spd = sp.drop_duplicates("ALPHA_3_CODE").set_index("ALPHA_3_CODE")

def _species_name(name_en, sci):
    """Аналог SQL-CASE: NAME_EN, при пустом — [SCIENTIFIC_NAME], иначе Unknown."""
    if isinstance(name_en, str) and name_en.strip():
        return name_en
    if isinstance(sci, str) and sci.strip():
        return "[" + sci.strip() + "]"
    return "Unknown"

sname = pd.Series(
    [_species_name(ne, sc) for ne, sc in zip(_spd.NAME_EN, _spd.SCIENTIFIC_NAME)],
    index=_spd.index,
)
cname = ctry.set_index("UN_CODE").NAME_EN
a_by_code = area.set_index("CODE").ID
p_by_child = rel_ima.drop_duplicates("CHILD_ID").set_index("CHILD_ID").PARENT_ID
wname = ima.set_index("ID").NAME_EN
ig_name = iss.set_index("ID").NAME_EN
ig_parent = iss_rel.drop_duplicates("CHILD_ID").set_index("CHILD_ID").PARENT_ID
# номер группы как в SQL: (CODE mod 10) + 10 * (CODE div 100)
code_num = pd.to_numeric(iss.set_index("ID").CODE)
ig_gnum = code_num % 10 + 10 * (code_num // 100)

for df in (q, v):
    df["_sid"] = df.SPECIES.map(sid)
    df["_inc"] = ~df._sid.isin(excl)
    df["_g2"] = df._sid.isin(g2)
    df["_g3"] = df._sid.isin(g3)
    df["_country"] = df.COUNTRY.map(cname).fillna("Unknown")
    df["_species"] = df.SPECIES.map(sname).fillna("Unknown")
    df["_waters"] = df.AREA.map(a_by_code).map(p_by_child).map(wname)
    df["_grp"] = df._sid.map(ig_parent).map(ig_name)
    df["_gnum"] = df._sid.map(ig_parent).map(ig_gnum)

qi, vi = q[q._inc], v[v._inc]

wb = load_workbook("aquaculture_report.xlsx", data_only=True)
mis = []

def num(x):
    return float(x) if isinstance(x, (int, float)) else 0.0

def cmp5(sheet, tag, got, exp):
    ev = [round(float(x), 2) for x in (exp.tolist() if hasattr(exp, "tolist") else list(exp))]
    for i, (g, e) in enumerate(zip(got, ev)):
        if abs(g - e) > 0.011:
            mis.append((sheet, tag, YEARS[i], g, e))

def row(ws, r, c0):
    return [num(ws.cell(row=r, column=c0 + i).value) for i in range(5)]

def series_or_zero(df_exp, key, cols=None):
    cols = cols or list(df_exp.columns)
    if key in df_exp.index:
        return df_exp.loc[key, cols]
    return [0.0] * len(cols)

# ---------- T1
ws = wb["Таблица 1"]
exp1 = qi.groupby("_country")[YC].sum()
seen = set()
for r in range(2, ws.max_row + 1):
    a = str(ws.cell(row=r, column=1).value or "")
    if not a: continue
    if a.startswith("ВСЕГО"):
        cmp5("T1", "TOTAL", row(ws, r, 2), exp1.sum())
    else:
        seen.add(a)
        cmp5("T1", a, row(ws, r, 2), series_or_zero(exp1, a))
missing = set(exp1[exp1.abs().sum(axis=1) > 1e-9].index) - seen
if missing: mis.append(("T1", "missing non-zero rows", sorted(missing)[:5], "", ""))

# ---------- T2
ws = wb["Таблица 2"]
exp2 = qi.groupby("_species")[YC].sum()
seen = set()
for r in range(2, ws.max_row + 1):
    a = str(ws.cell(row=r, column=1).value or "")
    if not a: continue
    if a.startswith("ВСЕГО"):
        cmp5("T2", "TOTAL", row(ws, r, 2), exp2.sum())
    else:
        seen.add(a)
        cmp5("T2", a, row(ws, r, 2), series_or_zero(exp2, a))
missing = set(exp2[exp2.abs().sum(axis=1) > 1e-9].index) - seen
if missing: mis.append(("T2", "missing non-zero rows", sorted(missing)[:5], "", ""))

# ---------- T3
ws = wb["Таблица 3"]
blocks = [("ВСЕГО", qi[qi._g2 | qi._g3], vi[vi._g2 | vi._g3]),
          ("РЫБА", qi[qi._g2], vi[vi._g2]),
          ("ВОДОРОСЛИ", qi[qi._g3], vi[vi._g3])]
oi = 0
nrows3 = 0
for r in range(2, ws.max_row + 1):
    a = str(ws.cell(row=r, column=1).value or "")
    b = str(ws.cell(row=r, column=2).value or "")
    if not a or a.strip() == "в том числе:":
        continue
    nrows3 += 1
    if a.startswith("ВСЕГО"): oi = 0
    elif a.startswith("РЫБА"): oi = 1
    elif a.startswith("ВОДОРОСЛИ"): oi = 2
    dq, dv = blocks[oi][1], blocks[oi][2]
    if a.startswith("     Морские"):
        dq, dv = dq[dq._waters == "Marine areas"], dv[dv._waters == "Marine areas"]
    elif a.startswith("     Внутренние"):
        dq, dv = dq[dq._waters == "Inland waters"], dv[dv._waters == "Inland waters"]
    src = dq if b == "К" else dv
    cmp5("T3", "%s|%s|%s" % (blocks[oi][0], a.strip(), b), row(ws, r, 3), src[YC].sum())
if nrows3 != 24:
    mis.append(("T3", "expected 24 data rows", nrows3, "", ""))

# ---------- T4
ws = wb["Таблица 4"]
expk = qi.groupby("_country")[YC].sum()
exps = vi.groupby("_country")[YC].sum()
cur = None
for r in range(2, ws.max_row + 1):
    a = str(ws.cell(row=r, column=1).value or "")
    b = str(ws.cell(row=r, column=2).value or "")
    if a: cur = "__T__" if a.startswith("ВСЕГО") else a
    E = expk if b == "К" else exps
    e = E.sum() if cur == "__T__" else series_or_zero(E, cur)
    cmp5("T4", "%s/%s" % (cur, b), row(ws, r, 3), e)

# ---------- T5
ws = wb["Таблица 5"]
expk = qi.groupby("_grp")[YC].sum()
exps = vi.groupby("_grp")[YC].sum()
for r in range(2, ws.max_row + 1):
    a = str(ws.cell(row=r, column=1).value or "")
    b = str(ws.cell(row=r, column=3).value or "")
    if not a: continue
    if a.startswith("ВСЕГО"):
        E = expk if b == "К" else exps
        cmp5("T5", "TOTAL/%s" % b, row(ws, r, 4), E.sum())
    else:
        E = expk if b == "К" else exps
        cmp5("T5", "%s/%s" % (a, b), row(ws, r, 4), series_or_zero(E, a))

# ---------- T6
ws = wb["Таблица 6"]
exp6 = qi.groupby(["_country", "_grp"])[YC].sum()
lv0 = exp6.index.get_level_values(0)
for r in range(2, ws.max_row + 1):
    a = str(ws.cell(row=r, column=1).value or "")
    g = str(ws.cell(row=r, column=2).value or "")
    if a.startswith("ВСЕГО"):
        cmp5("T6", "TOTAL", row(ws, r, 4), exp6.sum()); continue
    if not a: continue
    if g:
        key = (a, g)
        cmp5("T6", "%s/%s" % key, row(ws, r, 4),
             series_or_zero(exp6, key) if key in exp6.index else [0.0]*5)
    else:
        cmp5("T6", "SUBTOT/%s" % a, row(ws, r, 4),
             exp6[lv0 == a][YC].sum() if a in set(lv0) else [0.0]*5)

# ---------- T7
ws = wb["Таблица 7"]
exp7 = qi.groupby(["_country", "_species"])[YC].sum()
lv0 = exp7.index.get_level_values(0)
for r in range(2, ws.max_row + 1):
    a = str(ws.cell(row=r, column=1).value or "")
    s = str(ws.cell(row=r, column=2).value or "")
    if a.startswith("ВСЕГО"):
        cmp5("T7", "TOTAL", row(ws, r, 3), exp7.sum()); continue
    if not a: continue
    if s:
        key = (a, s)
        cmp5("T7", "%s/%s" % key, row(ws, r, 3),
             series_or_zero(exp7, key) if key in exp7.index else [0.0]*5)
    else:
        cmp5("T7", "SUBTOT/%s" % a, row(ws, r, 3),
             exp7[lv0 == a][YC].sum() if a in set(lv0) else [0.0]*5)

# ---------- T8
ws = wb["Таблица 8"]
exp8 = qi.groupby(["_grp", "_species", "ENVIRONMENT"])[YC].sum()
by_g8 = qi.groupby("_grp")[YC].sum()
by_env_all = qi.groupby("ENVIRONMENT")[YC].sum()
grand8 = by_env_all.sum()
cur_group = None
mode = "top"  # top -> until group totals appear; then rows
env_rows_seen = []
for r in range(2, ws.max_row + 1):
    g = str(ws.cell(row=r, column=1).value or "")
    s = str(ws.cell(row=r, column=3).value or "")
    env = str(ws.cell(row=r, column=4).value or "")
    if g.startswith("ВСЕГО"):
        cmp5("T8", "TOTAL", row(ws, r, 5), grand8); continue
    if not g and env and s == "" and cur_group is None and ":" not in env:
        # верхние итоги по средам (до первой группы)
        ekey = env.split(" ")[0]
        cmp5("T8", "ENVTOP/%s" % ekey, row(ws, r, 5),
             series_or_zero(by_env_all, ekey))
        continue
    if g and s == "ИТОГО ПО ГРУППЕ":
        cur_group = g
        cmp5("T8", "GTOT/%s" % g, row(ws, r, 5), series_or_zero(by_g8, g))
        continue
    if not g and env and cur_group is not None:
        ekey = env.split(" ")[0]
        sub = qi[qi._grp == cur_group]
        cmp5("T8", "GENV/%s/%s" % (cur_group, ekey), row(ws, r, 5),
             sub[sub.ENVIRONMENT == ekey][YC].sum())
        continue
    if g and s and env:
        key = (g, s, env)
        cmp5("T8", "%s/%s/%s" % key, row(ws, r, 5),
             series_or_zero(exp8, key) if key in exp8.index else [0.0]*5)

# ---------- T9
ws = wb["Таблица 9"]
exp9 = qi.groupby(["_grp", "_species"])[YC].sum()
by_g9 = qi.groupby("_grp")[YC].sum()
for r in range(2, ws.max_row + 1):
    g = str(ws.cell(row=r, column=1).value or "")
    s = str(ws.cell(row=r, column=3).value or "")
    if g.startswith("ВСЕГО"):
        cmp5("T9", "TOTAL", row(ws, r, 4), exp9.sum()); continue
    if not g: continue
    if not s:
        cmp5("T9", "GTOT/%s" % g, row(ws, r, 4), series_or_zero(by_g9, g))
    else:
        key = (g, s)
        cmp5("T9", "%s/%s" % key, row(ws, r, 4),
             series_or_zero(exp9, key) if key in exp9.index else [0.0]*5)

print("MISMATCHES:", len(mis))
for m in mis[:60]:
    print(m)
